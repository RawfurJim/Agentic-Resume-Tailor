import copy
import json
import os
import re

import docx
from docx.text.paragraph import Paragraph


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAIN_CV = os.path.join(ROOT, "Md_Rawfur_Monzur_Jim_CV.docx")
UPDATED_JSON = os.path.join(ROOT, "updated_cv.json")
OUTPUT_CV = os.path.join(ROOT, "Md_Rawfur_Monzur_Jim_CV_new.docx")


# ============================================================================
# HELPERS  (you never call these directly)
# ============================================================================

def is_bullet(p):
    """Is this paragraph a bullet? We look at the XML numbering, not the style
    name, because two of your bullets use the style 'No Spacing'."""
    return p._p.pPr is not None and p._p.pPr.numPr is not None


def get_section(doc, header):
    """All non-empty paragraphs between the ALL-CAPS line `header`
    (e.g. 'PROFILE') and the next ALL-CAPS line."""
    inside = False
    result = []
    for p in doc.paragraphs:
        text = p.text.strip()
        is_header = len(text) > 2 and text.isupper() and not is_bullet(p)
        if is_header:
            if inside:                      # reached the NEXT header -> done
                break
            inside = (text == header)       # start collecting after OUR header
        elif inside and text:
            result.append(p)
    return result


def read_text(p):
    """Paragraph -> text, with bold parts wrapped in **...**."""
    out = ""
    for run in p.runs:
        if run.bold and run.text.strip():
            out += "**" + run.text + "**"
        else:
            out += run.text
    return out.replace("****", "").strip()   # "**a****b**" -> "**ab**"


def set_text(p, new_text):
    """Replace the WHOLE text of one paragraph, keeping its look.
    Style / bullet / spacing are untouched because we only replace the runs."""
    # 1. remember the font settings of one existing run (prefer a non-bold one)
    template = None
    for run in p.runs:
        if run.text.strip() and not run.bold:
            template = run
            break
    if template is None and p.runs:
        template = p.runs[0]
    font_xml = None
    if template is not None and template._r.rPr is not None:
        font_xml = copy.deepcopy(template._r.rPr)

    # 2. delete every old run
    for run in p.runs:
        run._r.getparent().remove(run._r)

    # 3. add new runs: "aaa **bbb** ccc" -> "aaa " | "**bbb**" | " ccc"
    for chunk in re.split(r"(\*\*.+?\*\*)", new_text):
        if chunk == "":
            continue
        bold = chunk.startswith("**")
        run = p.add_run(chunk[2:-2] if bold else chunk)   # drop the ** at both ends
        if font_xml is not None:
            run._r.insert(0, copy.deepcopy(font_xml))    # same font as before
        if bold:
            run.bold = True


def make_count(paragraphs, n):
    """Return exactly n paragraphs. Need more? copy the last one.
    Have too many? delete from the end."""
    paragraphs = list(paragraphs)
    while len(paragraphs) < n:
        last = paragraphs[-1]
        new_el = copy.deepcopy(last._p)
        last._p.addnext(new_el)
        paragraphs.append(Paragraph(new_el, last._parent))
    while len(paragraphs) > n:
        gone = paragraphs.pop()
        gone._p.getparent().remove(gone._p)
    return paragraphs


def split_experience(doc):
    """The PROFESSIONAL EXPERIENCE section as
         job_header  (one paragraph)
         projects    [ {"title": <paragraph>, "bullets": [<paragraph>, ...]}, ... ]
    Rule: a bullet belongs to the last title seen; a non-bullet line is a new title."""
    paras = get_section(doc, "PROFESSIONAL EXPERIENCE")
    job_header = paras[0]
    projects = []
    for p in paras[1:]:
        if is_bullet(p):
            projects[-1]["bullets"].append(p)
        else:
            projects.append({"title": p, "bullets": []})
    return job_header, projects


# ============================================================================
# 1. GRAB EVERYTHING
# ============================================================================

def grab_all(doc):
    cv = {}

    cv["title"] = read_text(doc.paragraphs[1])                 # line under your name
    cv["profile"] = read_text(get_section(doc, "PROFILE")[0])

    cv["skills"] = []
    for p in get_section(doc, "KEY SKILLS"):
        label, items = p.text.split(":", 1)                    # "Label:  a, b, c"
        cv["skills"].append({"label": label.strip(), "items": items.strip()})

    job_header, projects = split_experience(doc)
    cv["role"] = job_header.text.split("—")[0].strip()         # "AI Engineer"
    cv["experience"] = []
    for proj in projects:
        cv["experience"].append({
            "title": proj["title"].text.strip(),
            "bullets": [read_text(b) for b in proj["bullets"]],
        })
    return cv


# ============================================================================
# 2. WRITE IT BACK  (updated_cv.json -> new docx)
# ============================================================================

class CVRewriter:
    """Puts the content of updated_cv.json into the main CV and saves a new docx."""

    def __init__(self, cv_path=MAIN_CV, json_path=UPDATED_JSON, output_path=OUTPUT_CV):
        self.cv_path = cv_path
        self.json_path = json_path
        self.output_path = output_path

    def rewrite(self):
        with open(self.json_path, encoding="utf-8") as f:
            cv = json.load(f)
        doc = docx.Document(self.cv_path)

        # title (line under the name)
        set_text(doc.paragraphs[1], cv["title"])

        # profile
        set_text(get_section(doc, "PROFILE")[0], cv["profile"])

        # key skills: "**Label:**  items"
        paras = make_count(get_section(doc, "KEY SKILLS"), len(cv["skills"]))
        for p, skill in zip(paras, cv["skills"]):
            set_text(p, f"**{skill['label']}:**  {skill['items']}")

        # professional experience
        job_header, projects = split_experience(doc)
        for run in job_header.runs:                       # "AI Engineer — Company\tDates"
            if "—" in run.text:
                run.text = cv["role"] + " " + run.text[run.text.index("—"):]
                break
        projects = self._make_projects(projects, len(cv["experience"]))
        for proj, data in zip(projects, cv["experience"]):
            set_text(proj["title"], data["title"])
            bullets = make_count(proj["bullets"], len(data["bullets"]))
            for b, text in zip(bullets, data["bullets"]):
                set_text(b, text)

        doc.save(self.output_path)
        return self.output_path

    @staticmethod
    def _make_projects(projects, n):
        """Like make_count, but for a whole project block (title + bullets)."""
        projects = list(projects)
        while len(projects) < n:
            last = projects[-1]
            anchor = last["bullets"][-1]._p
            title_el = copy.deepcopy(last["title"]._p)
            anchor.addnext(title_el)
            anchor = title_el
            bullets = []
            for b in last["bullets"]:
                el = copy.deepcopy(b._p)
                anchor.addnext(el)
                anchor = el
                bullets.append(Paragraph(el, b._parent))
            projects.append({"title": Paragraph(title_el, last["title"]._parent), "bullets": bullets})
        while len(projects) > n:
            gone = projects.pop()
            for p in [gone["title"], *gone["bullets"]]:
                p._p.getparent().remove(p._p)
        return projects


if __name__ == "__main__":
    print("Saved:", CVRewriter().rewrite())
