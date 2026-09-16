"""
Tests for the FastAPI backend (app/main.py).

Everything except the one test marked `integration` runs offline: the pipeline is
either never reached (validation / config errors) or replaced with a fake via
monkeypatch. The integration test makes a REAL DeepSeek call using the key in .env
(1-3 minutes, costs a few cents):

    pytest tests/test_api.py -m "not integration"    # fast, offline
    pytest tests/test_api.py -m integration -s       # real run
"""

import hashlib
import io

import docx
import pytest
from fastapi.testclient import TestClient

import app.main as main
from app.pipeline import MASTER_CV, PipelineError, PipelineResult

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

client = TestClient(main.app)


def _sha256(path: str) -> str:
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


# ---------------------------------------------------------------------------
# simple GET routes
# ---------------------------------------------------------------------------
def test_health():
    resp = client.get("/api/health")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}


def test_defaults():
    resp = client.get("/api/defaults")
    assert resp.status_code == 200
    body = resp.json()
    assert body["provider"] == "deepseek"
    assert body["deepseek"]["extract_model"] == "deepseek-flash"
    assert body["deepseek"]["rewrite_model"] == "deepseek-reasoner"
    assert body["gemini"]["extract_model"] == "gemini-2.5-flash"


def test_index_serves_html():
    resp = client.get("/")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/html")
    assert "CV Optimiser" in resp.text


# ---------------------------------------------------------------------------
# validation / config errors (pipeline never runs)
# ---------------------------------------------------------------------------
def test_blank_job_description_is_422():
    resp = client.post("/api/generate", json={"job_description": "   "})
    assert resp.status_code == 422
    assert "Job description is required" in resp.text


def test_missing_job_description_is_422():
    resp = client.post("/api/generate", json={"provider": "deepseek"})
    assert resp.status_code == 422


def test_gemini_without_key_is_400():
    resp = client.post("/api/generate", json={"job_description": "Some job", "provider": "gemini"})
    assert resp.status_code == 400
    assert "API key" in resp.json()["detail"]


def test_unknown_provider_is_422():
    resp = client.post("/api/generate", json={"job_description": "Some job", "provider": "openai"})
    assert resp.status_code == 422


def test_deepseek_without_any_key_is_400(monkeypatch):
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    resp = client.post("/api/generate", json={"job_description": "Some job"})
    assert resp.status_code == 400
    assert "DeepSeek API key" in resp.json()["detail"]


# ---------------------------------------------------------------------------
# error mapping and the download response, with run_pipeline faked out
# ---------------------------------------------------------------------------
def test_file_locked_maps_to_409(monkeypatch):
    def fake_run_pipeline(job_description, config):
        raise PipelineError("Close Md_Rawfur_Monzur_Jim_CV_new.docx and try again.", kind="file_locked")

    monkeypatch.setattr(main, "run_pipeline", fake_run_pipeline)
    resp = client.post("/api/generate", json={"job_description": "Some job"})
    assert resp.status_code == 409
    assert resp.json() == {"detail": "Close Md_Rawfur_Monzur_Jim_CV_new.docx and try again."}


@pytest.mark.parametrize(
    "kind, status",
    [("config", 400), ("llm", 502), ("parse", 502), ("other", 500)],
)
def test_other_pipeline_error_kinds(monkeypatch, kind, status):
    def fake_run_pipeline(job_description, config):
        raise PipelineError(f"boom {kind}", kind=kind)

    monkeypatch.setattr(main, "run_pipeline", fake_run_pipeline)
    resp = client.post("/api/generate", json={"job_description": "Some job"})
    assert resp.status_code == status
    assert resp.json() == {"detail": f"boom {kind}"}


def test_unexpected_exception_is_generic_500(monkeypatch):
    def fake_run_pipeline(job_description, config):
        raise RuntimeError("secret internals sk-abc")

    monkeypatch.setattr(main, "run_pipeline", fake_run_pipeline)
    resp = client.post("/api/generate", json={"job_description": "Some job", "api_key": "sk-abc"})
    assert resp.status_code == 500
    assert "sk-abc" not in resp.text            # nothing from the request or traceback leaks


def test_generate_returns_docx_download(monkeypatch):
    seen = {}

    def fake_run_pipeline(job_description, config):
        seen["job_description"] = job_description
        seen["config"] = config
        # Point at the master CV, read-only: we only need a real docx to stream back.
        return PipelineResult(docx_path=MASTER_CV, job_title="Senior AI Engineer / ML")

    monkeypatch.setattr(main, "run_pipeline", fake_run_pipeline)
    master_before = _sha256(MASTER_CV)

    resp = client.post("/api/generate", json={"job_description": "  Some job  "})

    assert resp.status_code == 200
    assert resp.headers["content-type"] == DOCX_MEDIA_TYPE
    disposition = resp.headers["content-disposition"]
    assert "attachment" in disposition
    assert "Md_Rawfur_Monzur_Jim_CV_Senior_AI_Engineer_ML.docx" in disposition
    assert resp.headers["x-job-title"] == "Senior AI Engineer / ML"

    document = docx.Document(io.BytesIO(resp.content))
    assert [p for p in document.paragraphs if p.text.strip()]

    # the validator stripped the whitespace and the defaults were applied
    assert seen["job_description"] == "Some job"
    assert seen["config"].provider == "deepseek"
    assert seen["config"].extract_model == "deepseek-flash"
    assert _sha256(MASTER_CV) == master_before


def test_ascii_header_value_strips_non_ascii_and_newlines():
    assert main.ascii_header_value("Ingénieur IA\r\nSenior") == "Ingenieur IA Senior"
    assert main.ascii_header_value("数据科学家") == ""


# ---------------------------------------------------------------------------
# the real thing: one DeepSeek run through the HTTP layer
# ---------------------------------------------------------------------------
SAMPLE_JD = """
Machine Learning Engineer (NLP) - Hybrid, London

Join our platform team to build and deploy NLP models that classify and summarise
customer support conversations. You will fine-tune transformer models, expose them
through Python REST services and monitor them in production on AWS.
Requirements: strong Python, PyTorch or similar, experience with LLM APIs and
prompt engineering, Docker, and a collaborative attitude in an Agile team.
"""


@pytest.mark.integration
def test_generate_end_to_end_with_deepseek():
    master_before = _sha256(MASTER_CV)

    resp = client.post("/api/generate", json={"job_description": SAMPLE_JD})
    print(f"\nstatus={resp.status_code} x-job-title={resp.headers.get('x-job-title')!r}")
    if resp.status_code != 200:
        print(resp.text)

    assert resp.status_code == 200
    assert resp.headers["content-type"] == DOCX_MEDIA_TYPE
    disposition = resp.headers["content-disposition"]
    assert "attachment" in disposition
    assert disposition.rstrip('"').endswith(".docx")
    assert resp.headers["x-job-title"].strip()

    document = docx.Document(io.BytesIO(resp.content))
    texts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    assert texts, "generated docx has no text paragraphs"

    assert _sha256(MASTER_CV) == master_before
